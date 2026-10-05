import datetime
import re
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)

book = Workbook()
values = book.active
values.title = "Values"
values["A1"] = "Text"
values["B1"] = "Açaí – São Paulo ✓"
values["A2"] = "Integer"
values["B2"] = 42
values["A3"] = "Decimal"
values["B3"] = -1234.5678
values["A4"] = "Boolean"
values["B4"] = True
values["A5"] = "Date"
values["B5"] = datetime.datetime(2026, 10, 4)
values["B5"].number_format = "yyyy-mm-dd"
values["A6"] = "Forced text"
values["B6"] = "0042"
values["B6"].number_format = "@"
values["A7"] = "Long text"
values["B7"] = "x" * 300
values["A8"] = "Comment"
values["B8"] = "noted"
values["B8"].comment = Comment("a note", "corpus")
values["A9"] = "Link"
values["B9"] = "Zekke"
values["B9"].hyperlink = "https://zekke.example/"

formulas = book.create_sheet("Data 2024")
for row, amount in enumerate([1200, 450.5, 80], start=1):
    formulas.cell(row=row, column=1, value=amount)
    formulas.cell(row=row, column=2, value="=A{0}*2".format(row))
formulas["A5"] = "=SUM(A1:A3)"
formulas["A6"] = "=$A$1+A2"
formulas["A7"] = "=Values!B2*2"
formulas["A8"] = "='Data 2024'!A1"
formulas["A9"] = "=IF(A1>1000,\"big\",\"small\")"
formulas["A10"] = "=_xlfn.XLOOKUP(450.5,A1:A3,B1:B3)"
formulas["A11"] = "=Total*1"
book.defined_names["Total"] = DefinedName("Total", attr_text="'Data 2024'!$A$5")

styles = book.create_sheet("Styles")
styles["A1"] = "bold"
styles["A1"].font = Font(bold=True)
styles["A2"] = "italic underline strike"
styles["A2"].font = Font(italic=True, underline="single", strike=True)
styles["A3"] = "red Arial 14"
styles["A3"].font = Font(name="Arial", size=14, color="FFFF0000")
styles["A4"] = "yellow fill"
styles["A4"].fill = PatternFill("solid", start_color="FFFFFF00", end_color="FFFFFF00")
styles["A5"] = "borders"
thin = Side(style="thin", color="FF000000")
thick = Side(style="thick", color="FF0000FF")
styles["A5"].border = Border(top=thin, bottom=thick, left=Side(style="dashed", color="FF00FF00"), right=Side(style="double", color="FF000000"))
styles["A6"] = "centered middle wrap"
styles["A6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
styles["A7"] = 1234.5
styles["A7"].number_format = '"R$" #,##0.00'
styles["A8"] = 0.256
styles["A8"].number_format = "0.0%"
styles["A9"] = 3.14159
styles["A9"].number_format = "0.00"
styles["A10"] = "right top"
styles["A10"].alignment = Alignment(horizontal="right", vertical="top")

layout = book.create_sheet("Layout")
layout["A1"] = "merged title"
layout.merge_cells("A1:C2")
layout.freeze_panes = "B4"
layout.column_dimensions["A"].width = 30
layout.column_dimensions["B"].width = 8
layout.column_dimensions["D"].hidden = True
layout.row_dimensions[5].height = 40
layout.row_dimensions[6].hidden = True
layout["A5"] = "tall row"
layout["A6"] = "hidden row"
layout["D1"] = "hidden column"
layout.sheet_properties.tabColor = "FF00AA00"

rules = book.create_sheet("Rules")
for row in range(1, 6):
    rules.cell(row=row, column=1, value=row * 10)
validation = DataValidation(type="list", formula1='"Paid,Pending,Overdue"', allow_blank=True)
rules.add_data_validation(validation)
validation.add("B1:B5")
rules.conditional_formatting.add("A1:A5", CellIsRule(operator="greaterThan", formula=["25"], fill=PatternFill("solid", start_color="FFFFC7CE", end_color="FFFFC7CE")))
rules.conditional_formatting.add("A1:A5", ColorScaleRule(start_type="min", start_color="FFFFFFFF", end_type="max", end_color="FF63BE7B"))

hidden = book.create_sheet("Hidden")
hidden["A1"] = "secret sheet"
hidden.sheet_state = "hidden"

plain = OUT / "corpus.plain.xlsx"
book.save(plain)

target = OUT / "corpus.xlsx"
with zipfile.ZipFile(plain) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as sink:
    for item in source.infolist():
        data = source.read(item.filename)
        if item.filename == "xl/worksheets/sheet2.xml":
            text = data.decode("utf-8")
            text = re.sub(r'<c r="B1"><f>A1\*2</f><v ?/?>(</v>)?</c>', '<c r="B1"><f t="shared" ref="B1:B3" si="0">A1*2</f><v>2400</v></c>', text)
            text = re.sub(r'<c r="B2"><f>A2\*2</f><v ?/?>(</v>)?</c>', '<c r="B2"><f t="shared" si="0"/><v>901</v></c>', text)
            text = re.sub(r'<c r="B3"><f>A3\*2</f><v ?/?>(</v>)?</c>', '<c r="B3"><f t="shared" si="0"/><v>160</v></c>', text)
            data = text.encode("utf-8")
        sink.writestr(item, data)
plain.unlink()

(OUT / "corpus.csv").write_text(
    'name,amount,note\r\n"Smith, John",1200.50,"line one\nline two"\r\nO\'Brien,-3,"says ""hi"""\r\n=1+1,0042,\r\n',
    encoding="utf-8",
)
(OUT / "corpus.tsv").write_text("name\tamount\nAçaí\t3.5\n", encoding="utf-8")
print("written", sorted(p.name for p in OUT.iterdir()))
